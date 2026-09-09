"""修改 Word 文档：查找替换、增删段落、编辑表格。"""

from __future__ import annotations

from typing import Dict, Optional

from docx import Document


def _replace_in_paragraph(paragraph, old: str, new: str) -> int:
    """在单个段落内替换文本（可跨多个 run），返回替换处数量。

    注意：替换后该段落原有 run 的格式会被重置为第一个 run 的格式。
    """
    if not old or old not in paragraph.text:
        return 0
    count = paragraph.text.count(old)
    runs = paragraph.runs
    if not runs:
        return 0
    replaced = paragraph.text.replace(old, new)
    for r in runs[1:]:
        r._element.getparent().remove(r._element)
    runs[0].text = replaced
    return count


def replace_text(doc: Document, old: str, new: str) -> int:
    """替换正文与表格中的所有匹配文本，返回替换处数量。

    不包含页眉 / 页脚（如需请在调用方单独处理）。
    """
    count = 0
    for p in doc.paragraphs:
        count += _replace_in_paragraph(p, old, new)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    count += _replace_in_paragraph(p, old, new)
    return count


def append_paragraph(doc: Document, text: str, style: str = None):
    """在文末追加一个段落。"""
    return doc.add_paragraph(text, style=style)


def prepend_paragraph(doc: Document, text: str, style: str = None):
    """在文首（第一个段落之前）插入一个段落。"""
    if doc.paragraphs:
        return doc.paragraphs[0].insert_paragraph_before(text, style=style)
    return doc.add_paragraph(text, style=style)


def remove_paragraphs_matching(doc: Document, text: str) -> int:
    """删除所有包含指定文本的正文段落，返回删除数量。"""
    removed = 0
    for p in list(doc.paragraphs):
        if text in p.text:
            p._element.getparent().remove(p._element)
            removed += 1
    return removed


def edit_docx(
    path: str,
    out_path: str,
    replacements: Optional[Dict[str, str]] = None,
    append: Optional[str] = None,
    prepend: Optional[str] = None,
    remove: Optional[str] = None,
) -> None:
    """加载 .docx，按顺序执行多项修改并另存。

    - ``replacements``：``{旧文本: 新文本}`` 查找替换
    - ``append`` / ``prepend``：追加 / 前插段落
    - ``remove``：删除所有包含该文本的段落
    """
    doc = Document(path)
    for old, new in (replacements or {}).items():
        replace_text(doc, old, new)
    if prepend:
        prepend_paragraph(doc, prepend)
    if append:
        append_paragraph(doc, append)
    if remove:
        remove_paragraphs_matching(doc, remove)
    doc.save(out_path)
