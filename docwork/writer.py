"""撰写 Word 文档：从 Markdown 文本或结构化数据创建 .docx。"""

from __future__ import annotations

import re
from typing import Any, Dict, List

from docx import Document
from docx.shared import Pt

from .utils import add_inline_runs, set_normal_font


def _new_document() -> Document:
    doc = Document()
    set_normal_font(doc)
    return doc


def _add_table(doc: Document, rows: List[List[str]]) -> None:
    """把行列数据（list of list of str）写入一个新表格。"""
    if not rows:
        return
    ncols = max(len(r) for r in rows)
    padded = [list(r) + [""] * (ncols - len(r)) for r in rows]
    table = doc.add_table(rows=len(padded), cols=ncols)
    table.style = "Table Grid"
    for ri, row in enumerate(padded):
        for ci, val in enumerate(row):
            table.rows[ri].cells[ci].text = val


def _add_code_block(doc: Document, code: str) -> None:
    p = doc.add_paragraph()
    run = p.add_run(code)
    run.font.name = "Consolas"
    p.paragraph_format.left_indent = Pt(12)


_SEPARATOR_CELL = re.compile(r":?-{2,}:?")


def markdown_to_docx(md_text: str, out_path: str) -> None:
    """从 Markdown 文本生成 .docx。

    支持：``#`` 标题、``-``/``*``/``+`` 无序列表、``1.`` 有序列表、
    表格（``| a | b |``）、``**加粗**``、``*斜体*``、`` `代码` ``、
    ````` ``` ````` 代码块。
    """
    doc = _new_document()
    lines = md_text.splitlines()
    i = 0
    in_code = False
    code_buf: List[str] = []
    table_buf: List[List[str]] = []

    def flush_table() -> None:
        nonlocal table_buf
        if not table_buf:
            return
        # 过滤 Markdown 表格的分隔行（如 | --- | --- |）
        rows = [
            r for r in table_buf
            if not all(_SEPARATOR_CELL.fullmatch(c) for c in r)
        ]
        _add_table(doc, rows)
        table_buf = []

    while i < len(lines):
        line = lines[i]
        s = line.strip()

        # 代码块
        if s.startswith("```"):
            if in_code:
                _add_code_block(doc, "\n".join(code_buf))
                code_buf = []
                in_code = False
            else:
                in_code = True
                code_buf = []
            i += 1
            continue
        if in_code:
            code_buf.append(line)
            i += 1
            continue

        # 表格行
        if s.startswith("|") and s.endswith("|") and len(s) > 1:
            table_buf.append([c.strip() for c in s.strip("|").split("|")])
            i += 1
            continue

        flush_table()

        if not s:
            i += 1
            continue

        # 标题
        if s.startswith("#"):
            level = len(s) - len(s.lstrip("#"))
            doc.add_heading(s.lstrip("#").strip(), level=min(level, 9))
        # 无序列表
        elif s.startswith(("- ", "* ", "+ ")):
            p = doc.add_paragraph(style="List Bullet")
            add_inline_runs(p, s[2:].strip())
        # 有序列表
        elif re.match(r"^\d+[.)]\s", s):
            text = re.sub(r"^\d+[.)]\s", "", s)
            p = doc.add_paragraph(style="List Number")
            add_inline_runs(p, text)
        # 普通段落
        else:
            p = doc.add_paragraph()
            add_inline_runs(p, s)
        i += 1

    flush_table()
    doc.save(out_path)


def _add_block(doc: Document, b: Dict[str, Any]) -> None:
    """把结构化块写入文档（与 read_docx 的输出结构保持一致）。"""
    t = b.get("type")
    if t == "heading":
        doc.add_heading(b.get("text", ""), level=b.get("level", 1))
    elif t == "table":
        _add_table(doc, b.get("rows", []))
    elif t == "list":
        style = "List Number" if b.get("ordered") else "List Bullet"
        for item in b.get("items", []):
            p = doc.add_paragraph(style=style)
            add_inline_runs(p, item)
    else:  # 默认按段落处理
        p = doc.add_paragraph()
        add_inline_runs(p, b.get("text", ""))


def structured_to_docx(data: Any, out_path: str, title: str = None) -> None:
    """从结构化数据创建 .docx。

    ``data`` 可为块列表（同 ``read_docx`` 输出），或 ``{"title": ..., "blocks": [...]}``。
    块类型支持 heading / paragraph / table / list。
    """
    doc = _new_document()
    if isinstance(data, dict):
        title = data.get("title", title)
        blocks = data.get("blocks", [])
    else:
        blocks = data

    if title:
        doc.add_heading(title, level=1)
    for b in blocks:
        _add_block(doc, b)
    doc.save(out_path)
